"""Load the Tata STRIVE partnerships workbook (Pipeline, MOU Status, Finance MIS sheets) into ObliGOT.

Real values from the sheets are used as they are: funders, Account Type (Tata / Non Tata / Government), locations,
account owners, Classification tiers (Finance MIS), pipeline deals with their dates, and MOUs with their dates and
signed-copy status. Only a partner's tier, group and location are simulated, and only when no sheet gives one;
each simulated value is listed in the partner's `simulated_fields` so the UI labels it. Nothing else is invented.

Usage (from backend/):
    python -I scripts/load_partnerships_workbook.py "<path to .xlsx>"          # dry run: prints what would change
    python -I scripts/load_partnerships_workbook.py "<path to .xlsx>" --write  # writes to the configured store

Re-running is safe: records get stable IDs and are overwritten, not duplicated.
"""
import hashlib
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from app.models import AgreementCreate, OpportunityCreate, PartnerCreate  # noqa: E402
from app.repository import list_records, put_record  # noqa: E402
from app.rules import business_today  # noqa: E402

GROUPS = {"tata": "tata_group", "non tata": "non_tata", "government": "government", "govt": "government"}
TIERS = {"platinum": "platinum", "gold": "gold", "silver": "silver", "bronze": "bronze"}
TIER_RANK = ["platinum", "gold", "silver", "bronze"]
# Simulation pools, used only when no sheet gives a value. Deterministic per funder name.
SIM_TIERS = ["platinum", "gold", "gold", "silver", "silver", "silver", "bronze", "bronze"]
SIM_GROUPS = ["tata_group", "non_tata", "government"]
SIM_CITIES = ["Mumbai", "Pune", "Bangalore", "Delhi", "Hyderabad", "Chennai", "Kolkata", "Ahmedabad"]
# Pipeline "Probability for approval" words → the deal's sanction %, stated so the mapping is visible.
PROBABILITY = {"high": 75, "medium": 50, "low": 25}


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:40] or "x"


def pick(name: str, pool: list[str], salt: str) -> str:
    return pool[int(hashlib.sha256(f"{salt}:{name}".encode()).hexdigest(), 16) % len(pool)]


def as_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def sheet(workbook, name: str, key: str) -> list[dict]:
    ws = next(w for w in workbook.worksheets if w.title.strip().lower() == name.lower())
    rows = list(ws.iter_rows(values_only=True))
    headers = [clean(c) for c in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if clean(dict(zip(headers, r)).get(key))]


def main(path: str, write: bool) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    pipeline = sheet(wb, "Pipeline", "Funder Name")
    mous = sheet(wb, "MOU Status", "Funder Name")
    finance = sheet(wb, "Finance MIS", "Funder Name")
    today = business_today().isoformat()

    # ---- partners: real facts gathered across sheets ----
    facts: dict[str, dict[str, Counter]] = {}
    def note(funder: str, field: str, value: str) -> None:
        if value:
            facts.setdefault(funder, {"group": Counter(), "location": Counter(), "tier": Counter(), "owner": Counter()})[field][value] += 1
    for row in pipeline + mous + finance:
        funder = clean(row.get("Funder Name"))
        facts.setdefault(funder, {"group": Counter(), "location": Counter(), "tier": Counter(), "owner": Counter()})
        note(funder, "group", GROUPS.get(clean(row.get("Account Type")).lower(), ""))
        note(funder, "owner", clean(row.get("Account Owner")))
        note(funder, "tier", TIERS.get(clean(row.get("Classification")).lower(), ""))
        place = clean(row.get("Location") or row.get("CenterCity"))
        note(funder, "location", "Multiple locations" if place.lower().startswith("multiple") else place)

    existing = {clean(p["name"]).lower(): p for p in list_records("partners")}
    partner_ids: dict[str, str] = {}
    plan: list[tuple[str, dict, str]] = []

    def resolve(name: str, f: dict[str, Counter] | None, base: dict) -> dict:
        simulated = []
        f = f or {"group": Counter(), "location": Counter(), "tier": Counter(), "owner": Counter()}
        # Values simulated on an earlier run are worked out again, so real data or a better rule replaces them.
        base = {k: v for k, v in base.items() if k not in base.get("simulated_fields", [])}
        tier = base.get("tier") or next((t for t in TIER_RANK if f["tier"].get(t)), None)  # highest real tier
        if not tier:
            tier, simulated = pick(name, SIM_TIERS, "tier"), simulated + ["tier"]
        group = base.get("group") or (f["group"].most_common(1)[0][0] if f["group"] else None)
        if not group and re.search(r"(?<![a-z])tata(?![a-z])", name, re.I):
            group = "tata_group"  # the organisation's own name says it is a Tata entity
        if not group:
            group, simulated = pick(name, SIM_GROUPS, "group"), simulated + ["group"]
        places = [p for p, _ in f["location"].most_common() if p != "Multiple locations"]
        location = base.get("location") or (places[0] if places else ("Multiple locations" if f["location"] else None))
        if not location:
            location, simulated = pick(name, SIM_CITIES, "city"), simulated + ["location"]
        kept = [s for s in base.get("simulated_fields", []) if s not in ("tier", "group", "location")]
        return {"tier": tier, "group": group, "location": location, "simulated_fields": kept + simulated}

    for funder, f in facts.items():
        current = existing.get(funder.lower())
        if current:
            partner_ids[funder] = current["id"]
            plan.append(("update partner", {**current, **resolve(funder, f, current), "updated_at": today}, current["id"]))
            continue
        owner = f["owner"].most_common(1)[0][0] if f["owner"] else "Unassigned"
        data = PartnerCreate(name=funder, category="CSR funder", status="active", owner=owner if len(owner) >= 2 else "Unassigned",
                             summary="Loaded from the partnerships workbook.").model_dump(mode="json")
        pid = f"partner-wb-{slug(funder)}"
        partner_ids[funder] = pid
        plan.append(("create partner", {**data, **resolve(funder, f, {}), "updated_at": today, "source_label": "workbook"}, pid))
    # Partners already in ObliGOT but not in the workbook (e.g. from uploaded MOUs): fill missing tier/group/location.
    for key, current in existing.items():
        if any(clean(fn).lower() == key for fn in facts):
            continue
        if current.get("simulated_fields") or not all(current.get(k) for k in ("tier", "group", "location")):
            plan.append(("update partner", {**current, **resolve(current["name"], None, current), "updated_at": today}, current["id"]))

    # ---- pipeline deals ----
    for n, row in enumerate(pipeline, 1):
        funder, prog = clean(row.get("Funder Name")), clean(row.get("ProgName"))
        final, proposal = clean(row.get("Final Status")).lower(), clean(row.get("Proposal Status")).lower()
        stage = "closed" if final == "decline" else "approval" if final == "approved" else \
            "qualification" if proposal == "in discussion" else "proposal"
        place = clean(row.get("Location"))
        revenue = row.get("Potential Revenue")
        deal = OpportunityCreate(
            partner_id=partner_ids[funder], title=prog if prog and prog.lower() != "skilling" else f"{funder} – {prog or 'Programme'}",
            stage=stage, owner=clean(row.get("Account Owner")) or "Unassigned", next_step=clean(row.get("Remark")),
            location=("Multiple locations" if place.lower().startswith("multiple") else place) or None,
            amount=float(revenue) if isinstance(revenue, (int, float)) else None,
            sanction_probability=PROBABILITY.get(clean(row.get("Probability for approval")).lower()),
            started_on=as_date(row.get("Opportunity Start Date")), last_followup_date=as_date(row.get("Last Follow Up Date")),
        ).model_dump(mode="json")
        plan.append(("create deal", {**deal, "updated_at": today, "source_label": "workbook"}, f"opportunity-wb-{n:03d}"))

    # ---- MOUs ----
    for n, row in enumerate(mous, 1):
        funder, prog = clean(row.get("Funder Name")), clean(row.get("ProgName"))
        kind, status = clean(row.get("Type of the agreement")).lower(), clean(row.get("MOU Status")).lower()
        signed = clean(row.get("MOU Signed Copy status")).lower()
        agreement = AgreementCreate(
            partner_id=partner_ids[funder],
            title=prog if prog.upper().startswith(funder.upper()) else f"{funder} – {prog}" if prog else funder,
            status="under_review" if kind == "proposal" else "expired" if "expired" in status else "active",
            start_date=as_date(row.get("MOU Start Date")), expiry_date=as_date(row.get("MOU End Date")),
            owner=clean(row.get("Account Owner")) or "Unassigned",
            signed_copy={"received": "received", "pending": "pending"}.get(signed, "unknown"),
            document_url=clean(row.get("MOU Link")) or None,
            programme_type=clean(row.get("Levers")) or None,
            total_target=float(row["Total Target"]) if isinstance(row.get("Total Target"), (int, float)) else None,
            total_budget=float(row["Total Budget"]) if isinstance(row.get("Total Budget"), (int, float)) else None,
            budget_currency="INR" if isinstance(row.get("Total Budget"), (int, float)) else None,
        ).model_dump(mode="json")
        plan.append(("create MOU", {**agreement, "source_label": "workbook"}, f"agreement-wb-{n:03d}"))

    for action, record, record_id in plan:
        sim = record.get("simulated_fields")
        extra = f" tier={record.get('tier')} group={record.get('group')} location={record.get('location')} simulated={sim}" if sim is not None else ""
        print(f"{action:15} {record_id:32} {record.get('name') or record.get('title')}{extra}")
    print(f"\n{len(plan)} writes planned.")
    if not write:
        print("Dry run only. Re-run with --write to save.")
        return
    for _, record, record_id in plan:
        put_record({"create partner": "partners", "update partner": "partners", "create deal": "opportunities",
                    "create MOU": "agreements"}[_], record, record_id)
    print("Saved.")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], "--write" in sys.argv[2:])
