"""Addendum comparison output remains document-scoped and is exported without approval."""
import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import app
from app.repository import get_record
from test_document_review import extracted, item

client = TestClient(app)


@pytest.fixture(autouse=True)
def setup(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path / "store"))
    monkeypatch.setattr(settings, "groq_api_key", "test-key")


def agreement(expiry="2026-11-30"):
    response = client.post("/api/agreements", json={
        "partner_id": "demo-003",
        "title": "Synthetic partnership agreement",
        "reference_number": "SYN-REV-002",
        "status": "active",
        "start_date": "2025-12-01",
        "expiry_date": expiry,
        "owner": "Test team",
    })
    assert response.status_code == 201, response.text
    return response.json()


def addendum_pages(expiry_text):
    return [
        "ADDENDUM No. 1 to MOU SYN-REV-002 between the parties (fictional). Signed on 1 October 2026.",
        f"The term of the MOU is extended until {expiry_text}. This addendum is effective from 1 October 2026.",
    ]


def addendum_output(expiry_iso, expiry_text):
    return {
        "fields": [
            item(field="reference_number", value_as_written="SYN-REV-002", page=1, quote="MOU SYN-REV-002"),
            item(field="execution_date", value_as_written="1 October 2026", normalized_value="2026-10-01",
                 page=1, quote="Signed on 1 October 2026"),
            item(field="expiry_date", value_as_written=expiry_text, normalized_value=expiry_iso,
                 page=2, quote=f"extended until {expiry_text}"),
            item(field="effective_date", value_as_written="1 October 2026", normalized_value="2026-10-01",
                 page=2, quote="effective from 1 October 2026"),
        ],
    }


def create_addendum(monkeypatch, expiry_iso="2027-03-31", expiry_text="31 March 2027"):
    return extracted(
        monkeypatch,
        output=addendum_output(expiry_iso, expiry_text),
        pages=addendum_pages(expiry_text),
        kind="addendum",
        partner_id="demo-003",
    )


def test_addendum_linking_remains_explicit(monkeypatch):
    target = agreement()
    doc = create_addendum(monkeypatch)
    candidates = client.get(f"/api/documents/{doc['id']}/agreement-candidates").json()
    assert any(candidate["agreement_id"] == target["id"] for candidate in candidates)
    assert get_record("documents", doc["id"])["agreement_id"] is None
    linked = client.patch(f"/api/documents/{doc['id']}/link", json={"agreement_id": target["id"]})
    assert linked.status_code == 200 and linked.json()["agreement_id"] == target["id"]


def test_addendum_changes_export_without_changing_operational_agreement(monkeypatch):
    target = agreement()
    doc = create_addendum(monkeypatch)
    assert client.patch(f"/api/documents/{doc['id']}/link", json={"agreement_id": target["id"]}).status_code == 200
    result = client.post(f"/api/documents/{doc['id']}/changes")
    assert result.status_code == 200
    expiry = next(change for change in result.json()["changes"] if change["field"] == "expiry_date")
    assert (expiry["prior_value"], expiry["proposed_value"], expiry["status"]) == (
        "2026-11-30", "2027-03-31", "proposed"
    )
    assert get_record("agreements", target["id"])["expiry_date"] == "2026-11-30"
    workbook_response = client.get(f"/api/documents/{doc['id']}/export/workbook")
    workbook = load_workbook(io.BytesIO(workbook_response.content))
    rows = list(workbook["Addendum_Changes"].iter_rows(values_only=True))
    assert any(row[2] == "expiry_date" and row[4] == "2027-03-31" and row[7] == "proposed" for row in rows[1:])


def test_addendum_decision_routes_are_removed(monkeypatch):
    target = agreement()
    doc = create_addendum(monkeypatch)
    client.patch(f"/api/documents/{doc['id']}/link", json={"agreement_id": target["id"]})
    result = client.post(f"/api/documents/{doc['id']}/changes").json()
    change_id = result["changes"][0]["id"]
    assert client.patch(f"/api/agreement-changes/{change_id}", json={"decision": "approve"}).status_code == 404
    assert client.post(f"/api/documents/{doc['id']}/approve", json={}).status_code == 404
