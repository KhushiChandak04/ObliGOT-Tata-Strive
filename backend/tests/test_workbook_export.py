"""Operational workbook export: contract conformance, agreement with the store, re-import."""
import io
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import app
from app.repository import list_records
from app.workbook import SHEETS
from test_document_review import extracted  # sibling module (pytest puts tests/ on sys.path)

client = TestClient(app)


@pytest.fixture(autouse=True)
def setup(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path / "store"))
    monkeypatch.setattr(settings, "groq_api_key", "test-key")


def export():
    response = client.get("/api/export/workbook")
    assert response.status_code == 200
    return response, load_workbook(io.BytesIO(response.content))


def table(wb, sheet):
    rows = list(wb[sheet].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def test_download_headers_and_sheets():
    response, wb = export()
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'attachment; filename="obligot-workbook-' in response.headers["content-disposition"]
    assert wb.sheetnames == list(SHEETS) + ["About"]
    assert wb.active.title == "Partners"
    assert client.get("/api/activities", params={"limit": 1}).json()[0]["event"] == "workbook_exported"


def test_every_sheet_follows_the_contract():
    _, wb = export()
    for sheet, columns in SHEETS.items():
        ws = wb[sheet]
        assert [c.value for c in ws[1]] == [c.name for c in columns], sheet
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
        for row in ws.iter_rows(min_row=2):
            for column, cell in zip(columns, row):
                value = cell.value
                if column.required and column.type != "bool":
                    assert value not in (None, ""), (sheet, column.name)
                if value is None:
                    continue
                if column.type == "date":
                    assert isinstance(value, datetime) and cell.number_format == "yyyy-mm-dd", (sheet, column.name)
                elif column.type in ("number", "integer"):
                    assert isinstance(value, (int, float)), (sheet, column.name)
                elif column.type == "bool":
                    assert isinstance(value, bool), (sheet, column.name)
                elif column.type == "enum":
                    assert value in column.allowed, (sheet, column.name, value)


def test_references_resolve_to_rows_in_the_referenced_sheet():
    _, wb = export()
    ids = {sheet: {r[SHEETS[sheet][0].name] for r in table(wb, sheet)} for sheet in SHEETS}
    for sheet, columns in SHEETS.items():
        for column in columns:
            if column.references:
                for row in table(wb, sheet):
                    if row[column.name] is not None:
                        assert row[column.name] in ids[column.references], (sheet, column.name, row[column.name])


def test_export_matches_the_store_and_leaves_missing_values_empty():
    _, wb = export()
    partners = table(wb, "Partners")
    assert {p["partner_id"] for p in partners} == {p["id"] for p in list_records("partners")}
    agreements = {a["agreement_id"]: a for a in table(wb, "Agreements")}
    for record in list_records("agreements"):
        row = agreements[record["id"]]
        assert (row["expiry_date"].date().isoformat() if row["expiry_date"] else None) == record.get("expiry_date")
        assert row["status"] == record["status"] and row["source_label"] == "demo"
    assert any(a["expiry_date"] is None for a in agreements.values())  # demo addendum has no expiry: left empty
    obligations = table(wb, "Obligations_Milestones")
    expected = len(list_records("milestones")) + len(list_records("reports")) + \
        sum(1 for f in list_records("finance") if f["record_type"] == "utilization_certificate")
    assert len(obligations) == expected
    assert not any(r.get("obligation_type") == "invoice" for r in obligations)  # invoices are not obligations here


def test_extracted_document_values_stay_out_of_operational_workbook(monkeypatch):
    doc = extracted(monkeypatch, partner_id="demo-003")
    _, wb = export()
    text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)
    assert doc["id"] not in text and doc["filename"] not in text


def test_partners_sheet_reimports_through_the_existing_preview():
    response, _ = export()
    preview = client.post("/api/import/preview", data={"entity_type": "partners"},
                          files={"file": ("obligot-workbook.xlsx", response.content, "application/octet-stream")}).json()
    assert preview["committed"] is False
    assert preview["headers"][:2] == ["partner_id", "name"]
    assert preview["valid_count"] == len(list_records("partners")) and preview["invalid_count"] == 0
