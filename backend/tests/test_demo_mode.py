"""Explicit synthetic demo mode never reads or writes the live Firestore project."""
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app import repository
from app.demo_dataset import build_demo_collections
from app.repository import DataStoreUnavailable

client = TestClient(app)


def _forbid_firestore():
    raise AssertionError("demo mode attempted to access Firestore")


def test_demo_mode_serves_linked_records_without_firestore(monkeypatch):
    monkeypatch.setattr(repository, "_DEMO_MODE", True)
    monkeypatch.setattr(repository, "_get_firestore", _forbid_firestore)

    documents = client.get("/api/documents")
    assert documents.status_code == 200
    assert len(documents.json()) == 6
    assert all(doc["source_label"] == "demo" for doc in documents.json())
    assert client.get("/api/health").json()["database"] == "demo-memory-fallback"
    assert client.get("/api/dashboard").json()["metrics"]["partners_total"] == 5
    assert len(client.get("/api/projects").json()) == 5
    assert len(client.get("/api/agreements").json()) == 6
    assert len(client.get("/api/finance").json()) == 4
    assert len(client.get("/api/reports").json()) == 3
    assert len(client.get("/api/milestones").json()) == 4
    assert len(client.get("/api/followups").json()) == 2
    assert client.get("/api/dashboard/summary").json()["partners"]["total"] == 5
    assert client.get("/api/documents", params={"kind": "addendum"}).json()[0]["id"] == "demo-document-005"
    detail = client.get("/api/documents/demo-document-001/extraction").json()
    assert detail["document"]["source_label"] == "demo"
    assert detail["extracted_fields"][0]["evidence"] is None

    partners = {record["id"] for record in client.get("/api/partners").json()}
    projects = {record["id"] for record in client.get("/api/projects").json()}
    agreements = {record["id"] for record in client.get("/api/agreements").json()}
    for project in client.get("/api/projects").json():
        assert project["partner_id"] in partners
    for agreement in client.get("/api/agreements").json():
        assert agreement["partner_id"] in partners
        assert agreement["project_id"] in projects
    for document in documents.json():
        assert document["partner_id"] in partners
        assert document["agreement_id"] in agreements
        assert document["page_count"] is None
        assert document["storage_ref"] is None
        assert "no source PDF" in document["warnings"][0]


def test_demo_writes_stay_in_memory_and_reset_on_backend_restart(monkeypatch):
    monkeypatch.setattr(repository, "_DEMO_MODE", True)
    monkeypatch.setattr(repository, "_get_firestore", _forbid_firestore)
    before = set(repository._COLLECTIONS["partners"])
    payload = {
        "name": "New Fictional Demo Partner", "category": "Training",
        "owner": "Synthetic Team", "summary": "Demo-only test record.",
    }

    created = client.post("/api/partners", json=payload)
    assert created.status_code == 201
    record_id = created.json()["id"]
    assert record_id in repository._DEMO_COLLECTIONS["partners"]
    assert set(repository._COLLECTIONS["partners"]) == before
    assert client.get("/api/partners").json()
    assert any(p["id"] == record_id for p in client.get("/api/partners").json())
    assert record_id not in build_demo_collections()["partners"]


def test_synthetic_document_and_operational_workbooks_are_labelled(monkeypatch):
    monkeypatch.setattr(repository, "_DEMO_MODE", True)
    monkeypatch.setattr(repository, "_get_firestore", _forbid_firestore)
    document_id = "demo-document-001"

    response = client.get(f"/api/documents/{document_id}/export/workbook")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    about = list(workbook["About"].values)
    assert any(row[0] == "Data source" and "SYNTHETIC DEMO" in row[1] for row in about)
    fields = list(workbook["Fields"].values)
    assert len(fields) == 2
    assert "Synthetic demo value" in str(fields[1][-1])

    operational = client.get("/api/export/workbook")
    assert operational.status_code == 200
    op_workbook = load_workbook(io.BytesIO(operational.content), read_only=True)
    assert any("SYNTHETIC DEMO" in str(row[0]) for row in op_workbook["About"].values)


def test_live_mode_failure_returns_503_without_falling_back_to_synthetic(monkeypatch):
    monkeypatch.setattr(repository, "_DEMO_MODE", False)
    monkeypatch.setattr(repository, "_TEST_MEMORY_STORE", False)

    def unavailable():
        raise DataStoreUnavailable("Firestore quota exceeded")

    monkeypatch.setattr(repository, "_get_firestore", unavailable)
    response = client.get("/api/documents")
    assert response.status_code == 503
    assert "live Firestore data store is unavailable" in response.json()["detail"]


def test_opted_in_firestore_read_failure_switches_to_demo_and_retries(monkeypatch):
    monkeypatch.setattr(repository, "_DEMO_MODE", False)
    monkeypatch.setattr(repository, "_AUTO_DEMO_FALLBACK", True)
    monkeypatch.setattr(repository, "_TEST_MEMORY_STORE", False)
    calls = []

    def unavailable():
        calls.append("init")
        raise DataStoreUnavailable("Firestore quota exceeded")

    monkeypatch.setattr(repository, "_get_firestore", unavailable)
    response = client.get("/api/documents")
    assert response.status_code == 200
    assert len(response.json()) == 6
    assert client.get("/api/documents").status_code == 200
    assert calls == ["init"]
    assert repository.is_demo_mode()
    assert response.headers["x-obligot-storage-mode"] == "demo-memory-fallback"


def test_opted_in_firestore_initialization_failure_switches_health_to_demo(monkeypatch):
    monkeypatch.setattr(repository, "_DEMO_MODE", False)
    monkeypatch.setattr(repository, "_AUTO_DEMO_FALLBACK", True)
    monkeypatch.setattr(repository, "_TEST_MEMORY_STORE", False)
    calls = []

    def unavailable():
        calls.append("init")
        raise DataStoreUnavailable("Firestore initialization failed")

    monkeypatch.setattr(repository, "_get_firestore", unavailable)
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["database"] == "demo-memory-fallback"
    assert response.headers["x-obligot-storage-mode"] == "demo-memory-fallback"
    assert calls == ["init"]


def test_opted_in_firestore_write_failure_keeps_write_in_demo_only(monkeypatch):
    monkeypatch.setattr(repository, "_DEMO_MODE", False)
    monkeypatch.setattr(repository, "_AUTO_DEMO_FALLBACK", True)
    monkeypatch.setattr(repository, "_TEST_MEMORY_STORE", False)

    def unavailable():
        raise DataStoreUnavailable("Firestore unavailable")

    monkeypatch.setattr(repository, "_get_firestore", unavailable)
    payload = {
        "name": "Fallback-only Fictional Partner", "category": "Training",
        "owner": "Synthetic Team",
    }
    response = client.post("/api/partners", json=payload)
    assert response.status_code == 201
    assert repository.is_demo_mode()
    assert response.headers["x-obligot-storage-mode"] == "demo-memory-fallback"
    assert any(record["name"] == payload["name"] for record in repository._DEMO_COLLECTIONS["partners"].values())
    assert all(getattr(record, "name", None) != payload["name"] for record in repository._COLLECTIONS["partners"].values())


def test_opted_in_firestore_delete_failure_retries_against_demo_only(monkeypatch):
    monkeypatch.setattr(repository, "_DEMO_MODE", False)
    monkeypatch.setattr(repository, "_AUTO_DEMO_FALLBACK", True)
    monkeypatch.setattr(repository, "_TEST_MEMORY_STORE", False)

    def unavailable():
        raise DataStoreUnavailable("Firestore unavailable")

    monkeypatch.setattr(repository, "_get_firestore", unavailable)
    demo_document_id = "demo-document-001"
    assert repository.delete_record("documents", demo_document_id)
    assert demo_document_id not in repository._DEMO_COLLECTIONS["documents"]
    assert repository.is_demo_mode()


def test_health_returns_error_for_unavailable_live_store_instead_of_demo_mode(monkeypatch):
    monkeypatch.setattr(repository, "_DEMO_MODE", False)
    monkeypatch.setattr(repository, "_TEST_MEMORY_STORE", False)

    def unavailable():
        raise DataStoreUnavailable("not configured")

    monkeypatch.setattr(repository, "_get_firestore", unavailable)

    response = client.get("/api/health")
    assert response.status_code == 503
    assert "live Firestore data store is unavailable" in response.json()["detail"]
