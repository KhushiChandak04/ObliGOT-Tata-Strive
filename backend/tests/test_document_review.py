"""Document-scoped extraction output and editable XLSX export; Groq is faked."""
import io
import json
import re

import pymupdf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import document_routes
from app.config import settings
from app.main import app
from app.repository import list_records

client = TestClient(app)

PAGES = [
    "SYNTHETIC MEMORANDUM OF UNDERSTANDING with Sample Learning Collective (fictional). Reference: SYN-REV-002.",
    "This MOU is effective from 1 December 2025 and remains valid until 30 November 2026.",
    "The partner shall submit an inception report by 30 September 2026. A utilization certificate is due by "
    "31 December 2026. Milestone: first cohort of 50 learners enrolled by 15 January 2027. The first tranche "
    "will be released on signing. Either party may terminate with 30 days written notice.",
]


def item(**kw):
    return kw


OUTPUT = {
    "fields": [
        item(field="partner_name", value_as_written="Sample Learning Collective (fictional)", page=1,
             quote="Sample Learning Collective (fictional)"),
        item(field="reference_number", value_as_written="SYN-REV-002", page=1, quote="Reference: SYN-REV-002"),
        item(field="effective_date", value_as_written="1 December 2025", normalized_value="2025-12-01", page=2,
             quote="effective from 1 December 2025"),
        item(field="expiry_date", value_as_written="30 November 2026", normalized_value="2026-11-30", page=2,
             quote="valid until 30 November 2026"),
    ],
    "obligations": [
        item(obligation_type="progress_report", title="Inception report", frequency="once", due_date="2026-09-30",
             page=3, quote="submit an inception report by 30 September 2026"),
        item(obligation_type="utilization_certificate", title="Utilization certificate", due_date="2026-12-31",
             page=3, quote="A utilization certificate is due by 31 December 2026"),
    ],
    "clauses": [
        item(clause_type="termination", summary="30 days' notice", page=3,
             quote="Either party may terminate with 30 days written notice.")
    ],
}


@pytest.fixture(autouse=True)
def setup(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "document_storage_dir", str(tmp_path / "store"))
    monkeypatch.setattr(settings, "groq_api_key", "test-key")


def make_pdf(pages):
    doc = pymupdf.open()
    for text in pages:
        doc.new_page().insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=10)
    data = doc.tobytes()
    doc.close()
    return data


def extracted(monkeypatch, output=OUTPUT, pages=PAGES, **form) -> dict:
    async def generate(prompt):
        present = {int(n) for n in re.findall(r"<<<PAGE (\d+)>>>", prompt)}
        return json.dumps({key: [row for row in rows if row["page"] in present] for key, rows in output.items()})

    monkeypatch.setattr(document_routes, "groq_generate", generate)
    doc = client.post(
        "/api/documents",
        data={"kind": "mou", **form},
        files={"file": ("synthetic-document.pdf", make_pdf(pages), "application/pdf")},
    ).json()
    response = client.post(f"/api/documents/{doc['id']}/extract")
    assert response.status_code == 202
    assert client.get(f"/api/documents/{doc['id']}").json()["status"] == "extracted"
    return doc


def proposals(doc_id):
    return client.get(f"/api/documents/{doc_id}/extraction").json()


def test_extraction_is_document_scoped_and_does_not_create_operational_records(monkeypatch):
    collections = ("partners", "agreements", "projects", "milestones", "reports", "finance")
    before = {collection: len(list_records(collection)) for collection in collections}
    doc = extracted(monkeypatch, partner_id="demo-003")
    output = proposals(doc["id"])
    assert output["document"]["status"] == "extracted"
    assert output["extracted_fields"] and output["obligations"] and output["clauses"]
    assert {collection: len(list_records(collection)) for collection in collections} == before


def test_document_approval_and_decision_routes_are_removed(monkeypatch):
    doc = extracted(monkeypatch)
    document_id = doc["id"]
    assert client.get(f"/api/documents/{document_id}/approval-check").status_code == 404
    assert client.post(f"/api/documents/{document_id}/approve", json={}).status_code == 404
    assert client.patch(
        f"/api/documents/{document_id}/proposals/extracted_fields/example",
        json={"decision": "approve"},
    ).status_code == 404
    assert client.patch(
        f"/api/documents/{document_id}/issues/example",
        json={"status": "resolved", "resolution": "checked"},
    ).status_code == 404


def test_document_workbook_is_editable_and_contains_extracted_values_and_evidence(monkeypatch):
    doc = extracted(monkeypatch)
    response = client.get(f"/api/documents/{doc['id']}/export/workbook")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"].endswith(f'obligot-{doc["id"]}-extraction.xlsx"')
    workbook = load_workbook(io.BytesIO(response.content))
    assert {"About", "Fields", "Clauses", "Financial_Terms", "Obligations", "Programme_Scope",
            "Tables", "Page_Readings", "Extraction_Notes", "Addendum_Changes"} <= set(workbook.sheetnames)
    fields = workbook["Fields"]
    row = next(row for row in fields.iter_rows(min_row=2, values_only=True) if row[1] == "partner_name")
    assert row[2] == "Sample Learning Collective (fictional)"
    assert row[4] == 1 and row[7] == "Sample Learning Collective (fictional)" and row[8] is True
    assert fields.protection.sheet is False and fields.auto_filter.ref.startswith("A1:")
    fields["C2"] = "Edited in Excel"
    saved = io.BytesIO()
    workbook.save(saved)
    saved.seek(0)
    assert load_workbook(saved)["Fields"]["C2"].value == "Edited in Excel"


def test_missing_document_workbook_returns_404():
    assert client.get("/api/documents/missing/export/workbook").status_code == 404
