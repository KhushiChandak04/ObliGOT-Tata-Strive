from datetime import date
from uuid import uuid4
from fastapi import APIRouter, HTTPException, Query, Request
from pydantic import ValidationError

from app.config import settings
from app.intelligence import analyze_context
from app.models import (
    ActionCreate, ActionItem, ActionUpdate, Agreement, AgreementCreate,
    AnalyzeRequest, AnalyzeResponse, FinanceCreate, FinanceRecord,
    FollowupCreate, FollowupUpdate, Opportunity, OpportunityCreate, Partner, PartnerCreate, PartnerUpdate,
    Project, ProjectCreate, ReportCreate, ReportRecord,
)
from app.repository import (
    firestore_enabled, get_record, list_records, patch_record, put_record, delete_record,
)
from app.rules import build_deadline_finding, business_today, scan_deadlines

router = APIRouter()

IMPORT_ENTITY_TYPES = {"agreements", "projects", "finance", "finance-records", "reports", "partners"}

@router.get("/health", tags=["system"])
def health() -> dict:
    return {"status": "ok", "app": "ObliGOT API", "environment": settings.environment,
            "database": "firestore" if firestore_enabled() else "demo-memory-fallback",
            "ai_configured": bool(settings.gemini_api_key), "auth_required": settings.require_firebase_auth}

@router.get("/dashboard", tags=["dashboard"])
def dashboard() -> dict:
    partners = list_records("partners")
    actions = list_records("actions")
    agreements = list_records("agreements")
    projects = list_records("projects")
    finance = list_records("finance")
    reports = list_records("reports")
    return {
        "metrics": {
            "partners_total": len(partners),
            "partners_active": sum(p.get("status") == "active" for p in partners),
            "actions_open": sum(a.get("status", "open") not in {"done", "deferred"} for a in actions),
            "agreements_active": sum(a.get("status") == "active" for a in agreements),
            "projects_active": sum(p.get("status") == "active" for p in projects),
            "reports_pending": sum(r.get("status") not in {"submitted", "accepted"} for r in reports),
            "finance_records": len(finance),
        },
        "findings": scan_deadlines({"agreements": agreements, "projects": projects, "finance": finance, "reports": reports}),
        "as_of": business_today().isoformat(),
        "storage_mode": "firestore" if firestore_enabled() else "demo-memory-fallback",
        "demo_data_notice": "Any record marked source_label=demo is illustrative, not verified.",
    }

@router.get("/partners", response_model=list[Partner], tags=["partners"])
def list_partners(q: str | None = Query(default=None, max_length=160)) -> list[Partner]:
    records = list_records("partners")
    if q:
        needle = q.casefold()
        records = [p for p in records if needle in p.get("name", "").casefold() or needle in p.get("category", "").casefold() or needle in p.get("owner", "").casefold()]
    return records

@router.post("/partners", response_model=Partner, status_code=201, tags=["partners"])
def create_partner(payload: PartnerCreate) -> dict:
    data = payload.model_dump(mode="json")
    data.update({"updated_at": business_today().isoformat(), "source_label": "user"})
    return put_record("partners", data)

@router.get("/partners/{partner_id}", response_model=Partner, tags=["partners"])
def get_partner(partner_id: str) -> dict:
    record = get_record("partners", partner_id)
    if not record:
        raise HTTPException(status_code=404, detail="Partner not found")
    return record

def _patch_validated(collection: str, record_id: str, changes: dict, model, label: str) -> dict:
    """Apply explicit changes, including nulls that clear optional fields, and validate the merged record."""
    current = get_record(collection, record_id)
    if not current:
        raise HTTPException(status_code=404, detail=f"{label} not found")
    merged = {**current, **changes}
    try:
        model.model_validate(merged)
    except ValidationError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    return put_record(collection, merged, record_id)

@router.patch("/partners/{partner_id}", response_model=Partner, tags=["partners"])
def update_partner(partner_id: str, payload: PartnerUpdate) -> dict:
    changes = payload.model_dump(exclude_unset=True, mode="json")
    changes["updated_at"] = business_today().isoformat()
    return _patch_validated("partners", partner_id, changes, Partner, "Partner")

# Collections that reference a partner, and the field holding the reference.
PARTNER_LINKS = {
    "agreements": "partner_id", "projects": "partner_id", "finance": "partner_id",
    "reports": "partner_id", "opportunities": "partner_id",
    "actions": "related_partner_id", "followups": "related_partner_id",
}

@router.delete("/partners/{partner_id}", status_code=204, tags=["partners"],
               responses={409: {"description": "Partner still has linked records"}})
def remove_partner(partner_id: str):
    if not get_record("partners", partner_id):
        raise HTTPException(status_code=404, detail="Partner not found")
    linked = {collection: count for collection, field in PARTNER_LINKS.items()
              if (count := sum(r.get(field) == partner_id for r in list_records(collection)))}
    if linked:
        summary = ", ".join(f"{collection}: {count}" for collection, count in linked.items())
        raise HTTPException(status_code=409, detail=f"Partner has linked records ({summary}). Reassign or delete them first.")
    if not delete_record("partners", partner_id):
        raise HTTPException(status_code=404, detail="Partner not found")
    return None

@router.get("/actions", response_model=list[ActionItem], tags=["actions"])
def list_actions(status: str | None = None, related_partner_id: str | None = None) -> list[dict]:
    records = list_records("actions")
    return [r for r in records
            if (not status or r.get("status", "open") == status)
            and (not related_partner_id or r.get("related_partner_id") == related_partner_id)]

@router.post("/actions", response_model=ActionItem, status_code=201, tags=["actions"])
def create_action(payload: ActionCreate) -> dict:
    data = payload.model_dump(mode="json")
    if data.get("related_partner_id") and not get_record("partners", data["related_partner_id"]):
        raise HTTPException(status_code=404, detail="Related partner not found")
    data.update({"status": "open", "source_label": "user"})
    return put_record("actions", data)

@router.get("/actions/{action_id}", response_model=ActionItem, tags=["actions"])
def get_action(action_id: str) -> dict:
    record = get_record("actions", action_id)
    if not record:
        raise HTTPException(status_code=404, detail="Action not found")
    return record

@router.patch("/actions/{action_id}", response_model=ActionItem, tags=["actions"])
def update_action(action_id: str, payload: ActionUpdate) -> dict:
    return _patch_validated("actions", action_id, payload.model_dump(exclude_unset=True, mode="json"), ActionItem, "Action")

@router.delete("/actions/{action_id}", status_code=204, tags=["actions"])
def delete_action(action_id: str):
    if not delete_record("actions", action_id):
        raise HTTPException(status_code=404, detail="Action not found")
    return None

@router.get("/agreements", response_model=list[Agreement], tags=["agreements"])
def list_agreements(partner_id: str | None = None) -> list[dict]:
    records = list_records("agreements")
    return [r for r in records if not partner_id or r.get("partner_id") == partner_id]

@router.post("/agreements", response_model=Agreement, status_code=201, tags=["agreements"])
def create_agreement(payload: AgreementCreate) -> dict:
    if not get_record("partners", payload.partner_id):
        raise HTTPException(status_code=404, detail="Partner not found")
    return put_record("agreements", {**payload.model_dump(mode="json"), "source_label": "user"})

async def _read_json_object(request: Request) -> dict:
    try:
        body = await request.json()
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="Request body must be valid JSON") from exc
    if not isinstance(body, dict):
        raise HTTPException(status_code=422, detail="Request body must be a JSON object")
    return body

async def _update_typed_record(collection: str, record_id: str, request: Request, schema, label: str, extra: dict | None = None) -> dict:
    changes = await _read_json_object(request)
    if not isinstance(changes, dict):
        raise HTTPException(status_code=422, detail="Request body must be a JSON object")
    current = get_record(collection, record_id)
    if not current:
        raise HTTPException(status_code=404, detail=f"{label} not found")
    merged = {k: v for k, v in current.items() if k not in {"id", "source_label"}}
    merged.update(changes)
    try:
        validated = schema.model_validate(merged)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    data = validated.model_dump(mode="json")
    if data.get("partner_id") and not get_record("partners", data["partner_id"]):
        raise HTTPException(status_code=404, detail="Partner not found")
    data["source_label"] = current.get("source_label", "user")
    data.update(extra or {})
    return put_record(collection, data, record_id)

@router.get("/agreements/{agreement_id}", response_model=Agreement, tags=["agreements"])
def get_agreement(agreement_id: str) -> dict:
    record = get_record("agreements", agreement_id)
    if not record:
        raise HTTPException(status_code=404, detail="Agreement not found")
    return record

@router.patch("/agreements/{agreement_id}", response_model=Agreement, tags=["agreements"])
async def update_agreement(agreement_id: str, request: Request) -> dict:
    return await _update_typed_record("agreements", agreement_id, request, AgreementCreate, "Agreement")

@router.delete("/agreements/{agreement_id}", status_code=204, tags=["agreements"])
def delete_agreement(agreement_id: str):
    if not delete_record("agreements", agreement_id):
        raise HTTPException(status_code=404, detail="Agreement not found")
    return None

@router.get("/projects", response_model=list[Project], tags=["projects"])
def list_projects(partner_id: str | None = None) -> list[dict]:
    records = list_records("projects")
    return [r for r in records if not partner_id or r.get("partner_id") == partner_id]

@router.post("/projects", response_model=Project, status_code=201, tags=["projects"])
def create_project(payload: ProjectCreate) -> dict:
    if payload.partner_id and not get_record("partners", payload.partner_id):
        raise HTTPException(status_code=404, detail="Partner not found")
    data = payload.model_dump(mode="json")
    data["source_label"] = "user"
    return put_record("projects", data)

@router.get("/projects/{project_id}", response_model=Project, tags=["projects"])
def get_project(project_id: str) -> dict:
    record = get_record("projects", project_id)
    if not record:
        raise HTTPException(status_code=404, detail="Project not found")
    return record

@router.patch("/projects/{project_id}", response_model=Project, tags=["projects"])
async def update_project(project_id: str, request: Request) -> dict:
    changes = await _read_json_object(request)
    allowed = {"partner_id", "name", "status", "owner", "start_date", "end_date", "milestone_summary"}
    unknown = set(changes) - allowed
    if unknown:
        raise HTTPException(status_code=422, detail=f"Unsupported project fields: {', '.join(sorted(unknown))}")
    if "partner_id" in changes and changes["partner_id"] and not get_record("partners", changes["partner_id"]):
        raise HTTPException(status_code=404, detail="Partner not found")
    current = get_record("projects", project_id)
    if not current:
        raise HTTPException(status_code=404, detail="Project not found")
    merged = {k: v for k, v in current.items() if k in allowed}
    merged.update(changes)
    try:
        validated = ProjectCreate.model_validate(merged)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    return put_record("projects", {**validated.model_dump(mode="json"), "source_label": current.get("source_label", "user")}, project_id)

@router.delete("/projects/{project_id}", status_code=204, tags=["projects"])
def delete_project(project_id: str):
    if not delete_record("projects", project_id):
        raise HTTPException(status_code=404, detail="Project not found")
    return None

@router.get("/finance/{finance_id}", response_model=FinanceRecord, tags=["finance"])
def get_finance_record(finance_id: str) -> dict:
    record = get_record("finance", finance_id)
    if not record:
        raise HTTPException(status_code=404, detail="Finance record not found")
    return record

@router.patch("/finance/{finance_id}", response_model=FinanceRecord, tags=["finance"])
async def update_finance_record(finance_id: str, request: Request) -> dict:
    return await _update_typed_record("finance", finance_id, request, FinanceCreate, "Finance record")

@router.delete("/finance/{finance_id}", status_code=204, tags=["finance"])
def delete_finance_record(finance_id: str):
    if not delete_record("finance", finance_id):
        raise HTTPException(status_code=404, detail="Finance record not found")
    return None

@router.get("/finance", response_model=list[FinanceRecord], tags=["finance"])
def list_finance(partner_id: str | None = None) -> list[dict]:
    records = list_records("finance")
    return [r for r in records if not partner_id or r.get("partner_id") == partner_id]

@router.post("/finance", response_model=FinanceRecord, status_code=201, tags=["finance"])
def create_finance(payload: FinanceCreate) -> dict:
    if not get_record("partners", payload.partner_id):
        raise HTTPException(status_code=404, detail="Partner not found")
    return put_record("finance", {**payload.model_dump(mode="json"), "source_label": "user"})

@router.get("/reports/{report_id}", response_model=ReportRecord, tags=["reports"])
def get_report(report_id: str) -> dict:
    record = get_record("reports", report_id)
    if not record:
        raise HTTPException(status_code=404, detail="Report not found")
    return record

@router.patch("/reports/{report_id}", response_model=ReportRecord, tags=["reports"])
async def update_report(report_id: str, request: Request) -> dict:
    return await _update_typed_record("reports", report_id, request, ReportCreate, "Report")

@router.delete("/reports/{report_id}", status_code=204, tags=["reports"])
def delete_report(report_id: str):
    if not delete_record("reports", report_id):
        raise HTTPException(status_code=404, detail="Report not found")
    return None

@router.get("/reports", response_model=list[ReportRecord], tags=["reports"])
def list_reports(partner_id: str | None = None) -> list[dict]:
    records = list_records("reports")
    return [r for r in records if not partner_id or r.get("partner_id") == partner_id]

@router.post("/reports", response_model=ReportRecord, status_code=201, tags=["reports"])
def create_report(payload: ReportCreate) -> dict:
    if not get_record("partners", payload.partner_id):
        raise HTTPException(status_code=404, detail="Partner not found")
    return put_record("reports", {**payload.model_dump(mode="json"), "source_label": "user"})

@router.get("/opportunities", response_model=list[Opportunity], tags=["pipeline"])
def list_opportunities(partner_id: str | None = None, stage: str | None = None) -> list[dict]:
    records = list_records("opportunities")
    return [r for r in records
            if (not partner_id or r.get("partner_id") == partner_id)
            and (not stage or r.get("stage") == stage)]

@router.post("/opportunities", response_model=Opportunity, status_code=201, tags=["pipeline"])
def create_opportunity(payload: OpportunityCreate) -> dict:
    if not get_record("partners", payload.partner_id):
        raise HTTPException(status_code=404, detail="Partner not found")
    return put_record("opportunities", {**payload.model_dump(mode="json"), "updated_at": business_today().isoformat(), "source_label": "user"})

@router.get("/opportunities/{opportunity_id}", response_model=Opportunity, tags=["pipeline"])
def get_opportunity(opportunity_id: str) -> dict:
    record = get_record("opportunities", opportunity_id)
    if not record:
        raise HTTPException(status_code=404, detail="Opportunity not found")
    return record

@router.patch("/opportunities/{opportunity_id}", response_model=Opportunity, tags=["pipeline"])
async def update_opportunity(opportunity_id: str, request: Request) -> dict:
    return await _update_typed_record("opportunities", opportunity_id, request, OpportunityCreate, "Opportunity",
                                      extra={"updated_at": business_today().isoformat()})

@router.delete("/opportunities/{opportunity_id}", status_code=204, tags=["pipeline"])
def delete_opportunity(opportunity_id: str):
    if not delete_record("opportunities", opportunity_id):
        raise HTTPException(status_code=404, detail="Opportunity not found")
    return None

@router.get("/followups", tags=["follow-ups"])
def list_followups(status: str | None = None, related_partner_id: str | None = None) -> list[dict]:
    records = list_records("followups")
    return [r for r in records
            if (not status or r.get("status", "open") == status)
            and (not related_partner_id or r.get("related_partner_id") == related_partner_id)]

@router.get("/followups/{followup_id}", tags=["follow-ups"])
def get_followup(followup_id: str) -> dict:
    record = get_record("followups", followup_id)
    if not record:
        raise HTTPException(status_code=404, detail="Follow-up not found")
    return record

@router.delete("/followups/{followup_id}", status_code=204, tags=["follow-ups"])
def delete_followup(followup_id: str):
    if not delete_record("followups", followup_id):
        raise HTTPException(status_code=404, detail="Follow-up not found")
    return None

@router.post("/intelligence/analyze", response_model=AnalyzeResponse, tags=["intelligence"])
async def analyze(payload: AnalyzeRequest) -> AnalyzeResponse:
    partner_name = None
    context = payload.context
    if payload.partner_id:
        partner = get_record("partners", payload.partner_id)
        if not partner:
            raise HTTPException(status_code=404, detail="Partner not found")
        partner_name = partner.get("name")
        context = f"Partner summary: {partner.get('summary', '')}\nAdditional context: {context}"
    return await analyze_context(context, partner_name)

@router.post("/followups", status_code=201, tags=["follow-ups"])
def create_followup(payload: FollowupCreate) -> dict:
    if payload.related_partner_id and not get_record("partners", payload.related_partner_id):
        raise HTTPException(status_code=404, detail="Related partner not found")
    data = {**payload.model_dump(mode="json"), "status": "open", "source_label": "user"}
    return put_record("followups", data)

@router.patch("/followups/{followup_id}", tags=["follow-ups"])
def update_followup(followup_id: str, payload: FollowupUpdate) -> dict:
    record = patch_record("followups", followup_id, payload.model_dump(mode="json"))
    if not record:
        raise HTTPException(status_code=404, detail="Follow-up not found")
    return record

@router.get("/rules/deadline-preview", tags=["rules"])
def deadline_preview(record_id: str = "demo-record", record_type: str = "report", label: str = "Quarterly report", due_date: date | None = None) -> dict:
    return {"finding": build_deadline_finding(record_id=record_id, record_type=record_type, label=label, due_date=due_date),
            "note": "Preview only; use verified record dates for operational use."}

@router.post("/import/preview", tags=["import"])
async def import_preview(request: Request) -> dict:
    """Validate JSON row batches or preview CSV/XLSX uploads without writing records."""
    import csv
    import io
    content_type = request.headers.get("content-type", "").lower()
    filename = None
    headers: list[str] = []
    if "application/json" in content_type:
        payload = await _read_json_object(request)
        entity_type = payload.get("entity_type")
        if entity_type not in IMPORT_ENTITY_TYPES:
            raise HTTPException(status_code=422, detail="Unsupported entity_type")
        rows = payload.get("rows")
        if not isinstance(rows, list) or not rows or len(rows) > 500:
            raise HTTPException(status_code=422, detail="rows must contain between 1 and 500 records")
    elif "multipart/form-data" in content_type:
        form = await request.form()
        entity_type = form.get("entity_type", "projects")
        if entity_type not in IMPORT_ENTITY_TYPES:
            raise HTTPException(status_code=422, detail="Unsupported entity_type")
        upload = form.get("file")
        if upload is None or not hasattr(upload, "read"):
            raise HTTPException(status_code=422, detail="A file field is required")
        filename = getattr(upload, "filename", None)
        raw = await upload.read()
        if len(raw) > 5_000_000:
            raise HTTPException(status_code=413, detail="File is larger than the 5 MB preview limit")
        try:
            if (filename or "").lower().endswith(".csv"):
                decoded = raw.decode("utf-8-sig")
                reader = csv.DictReader(io.StringIO(decoded))
                headers = list(reader.fieldnames or [])
                rows = list(reader)[:500]
            elif (filename or "").lower().endswith((".xlsx", ".xlsm")):
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                sheet = workbook.active
                iterator = sheet.iter_rows(values_only=True)
                first = next(iterator, None)
                headers = [str(v).strip() if v is not None else "" for v in (first or [])]
                rows = [dict(zip(headers, [str(v) if v is not None else "" for v in row])) for row in iterator]
                rows = rows[:500]
                workbook.close()
            else:
                raise HTTPException(status_code=415, detail="Upload a .csv or .xlsx file")
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=400, detail="CSV must be UTF-8 encoded") from exc
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=400, detail="Could not read this file; check its format") from exc
    else:
        raise HTTPException(status_code=415, detail="Send JSON rows or upload a CSV/XLSX file")

    seen: set[str] = set()
    valid_rows: list[dict] = []
    issues: list[dict] = []
    for idx, row in enumerate(rows, 1):
        if not isinstance(row, dict):
            row = {}
        name = str(row.get("name", "") or "").strip()
        errors: list[str] = []
        if not name:
            errors.append("Missing required field: name")
        elif name.casefold() in seen:
            errors.append("Duplicate name in this batch")
        if name:
            seen.add(name.casefold())
        amount = row.get("amount")
        if amount not in (None, ""):
            try:
                if float(amount) < 0:
                    errors.append("Amount cannot be negative")
            except (TypeError, ValueError):
                errors.append("Amount must be numeric")
        item = {"row": idx, "name": name or None, "valid": not errors, "issues": errors}
        (valid_rows if not errors else issues).append(item)

    return {
        "entity_type": entity_type,
        "filename": filename,
        "headers": headers,
        "preview_rows": rows[:100],
        "rows_previewed": min(len(rows), 100),
        "total_rows_detected": len(rows),
        "total_rows": len(rows),
        "valid_count": len(valid_rows),
        "invalid_count": len(issues),
        "valid_rows": valid_rows,
        "issues": issues,
        "committed": False,
        "warning": "Preview only. Review mappings and validation issues before importing; no records were written.",
    }
