"""Build an editable, document-scoped workbook of extraction output and its evidence."""
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.repository import is_demo_mode, list_records

HEADER_FILL = PatternFill("solid", fgColor="1F4D3F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _text(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (list, dict)):
        return str(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _evidence(record: dict) -> tuple[Any, ...]:
    evidence = record.get("evidence") or {}
    return (
        evidence.get("page"),
        evidence.get("printed_page"),
        evidence.get("section"),
        evidence.get("quote"),
        evidence.get("verified"),
        evidence.get("verification_error"),
        evidence.get("basis"),
    )


def _append_sheet(wb: Workbook, name: str, headers: list[str], rows: list[list[Any]]):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        cell.data_type = "s"
    for row in rows:
        ws.append([_text(value) for value in row])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        longest = max((len(str(cell.value)) for cell in column if cell.value is not None), default=0)
        ws.column_dimensions[letter].width = min(max(longest + 2, 12), 60)
    return ws


def build_document_workbook(document: dict, proposals: dict[str, list[dict]]) -> bytes:
    """Return a safe-to-edit XLSX containing extracted values, page evidence, and extraction notes."""
    wb = Workbook()
    wb.remove(wb.active)
    _append_sheet(wb, "About", ["Item", "Value"], [
        ["Filename", document["filename"]],
        ["Data source", "SYNTHETIC DEMO — NOT CONNECTED TO LIVE FIRESTORE RECORDS" if is_demo_mode() else "Document record"],
        ["Document type", document["kind"]],
        ["Extraction status", document["status"]],
        ["Extraction mode", (document.get("extraction") or {}).get("mode")],
        ["Extraction note", (document.get("extraction") or {}).get("note")],
        ["Uploaded at", document.get("uploaded_at")],
        ["Partner", document.get("partner_id")],
        ["Agreement", document.get("agreement_id")],
        ["How to use", "This workbook is editable. Changes made here are not saved back to ObliGOT."],
        ["Review note", "AI-extracted values are not verified business records. Compare source pages and evidence before relying on them."],
        ["Data impact", "This document-scoped export does not create or update operational partner, agreement, project, finance, or obligation records."],
        ["Synthetic data warning", "No source PDF, actual extraction, or source evidence exists for this demo record." if document.get("source_label") == "demo" else None],
    ])

    fields = proposals.get("extracted_fields", [])
    _append_sheet(wb, "Fields", [
        "Target", "Field", "Extracted value", "As written", "Source page", "Printed page",
        "Section", "Source quote", "Evidence found", "Evidence error", "Evidence basis", "Warnings",
    ], [[r.get("target"), r.get("field"), r.get("proposed_value"), r.get("verbatim"), *_evidence(r), "\n".join(r.get("warnings") or [])]
        for r in fields])

    clauses = proposals.get("clauses", [])
    _append_sheet(wb, "Clauses", [
        "Clause type", "Heading", "Summary", "Verbatim text", "Needs legal review",
        "Source page", "Printed page", "Section", "Source quote", "Evidence found", "Evidence error", "Evidence basis", "Warnings",
    ], [[r.get("clause_type"), r.get("heading"), r.get("summary"), r.get("verbatim"), r.get("needs_legal_review"),
         *_evidence(r), "\n".join(r.get("warnings") or [])] for r in clauses])

    terms = proposals.get("financial_terms", [])
    _append_sheet(wb, "Financial_Terms", [
        "Term type", "Description", "Amount", "Currency", "Basis", "Financial year", "Payer", "Due date", "Condition",
        "Source page", "Printed page", "Section", "Source quote", "Evidence found", "Evidence error", "Evidence basis", "Warnings",
    ], [[r.get("term_type"), r.get("description"), r.get("amount"), r.get("currency"), r.get("basis"),
         r.get("financial_year"), r.get("payer"), r.get("due_date"), r.get("condition"), *_evidence(r),
         "\n".join(r.get("warnings") or [])] for r in terms])

    obligations = proposals.get("obligations", [])
    _append_sheet(wb, "Obligations", [
        "Obligation type", "Title", "Frequency", "Due date", "Deadline type", "Deadline text", "Offset value",
        "Offset unit", "Trigger", "Trigger event", "Recurrence", "Owner", "Source page", "Printed page",
        "Section", "Source quote", "Evidence found", "Evidence error", "Evidence basis", "Warnings",
    ], [[
        r.get("obligation_type"), r.get("title"), r.get("frequency"), r.get("due_date"),
        (r.get("deadline") or {}).get("deadline_type"), (r.get("deadline") or {}).get("text"),
        (r.get("deadline") or {}).get("offset_value"), (r.get("deadline") or {}).get("offset_unit"),
        (r.get("deadline") or {}).get("trigger"), (r.get("deadline") or {}).get("trigger_event"),
        (r.get("deadline") or {}).get("recurrence"), r.get("owner"), *_evidence(r),
        "\n".join(r.get("warnings") or []),
    ] for r in obligations])

    scope = proposals.get("programme_scope", [])
    _append_sheet(wb, "Programme_Scope", [
        "Dimension", "Value", "Target value", "Target unit", "Location", "Period", "Source page", "Printed page",
        "Section", "Source quote", "Evidence found", "Evidence error", "Evidence basis", "Warnings",
    ], [[r.get("dimension"), r.get("value"), r.get("target_value"), r.get("target_unit"), r.get("location"),
         r.get("period"), *_evidence(r), "\n".join(r.get("warnings") or [])] for r in scope])

    tables = sorted(proposals.get("tables", []), key=lambda record: record["id"])
    index_rows = []
    for number, table in enumerate(tables, start=1):
        sheet_name = f"Table_{number}"
        index_rows.append([table.get("title"), table.get("kind"), table.get("section"),
                           ", ".join(map(str, table.get("pages") or [])), table.get("currency"),
                           "\n".join(table.get("unit_notes") or []), "\n".join(table.get("notes") or []), sheet_name])
        headers = ["Source page", *[str(header) or f"Column {i}" for i, header in enumerate(table.get("headers") or [], start=1)]]
        rows = [[row.get("page"), *(row.get("cells") or [])] for row in table.get("rows") or []]
        _append_sheet(wb, sheet_name, headers, rows)
    _append_sheet(wb, "Tables", ["Title", "Kind", "Section", "Pages", "Currency", "Unit notes", "Notes", "Data sheet"], index_rows)

    pages = sorted((p for p in list_records("document_pages") if p.get("document_id") == document["id"]),
                   key=lambda page: page["page_number"])
    _append_sheet(wb, "Page_Readings", [
        "Physical page", "Printed page", "Reading method", "Page type", "Legibility", "Transcript",
        "Unreadable regions", "Handwriting or stamp", "Has table",
    ], [[p.get("page_number"), p.get("printed_page_number"), p.get("method"), p.get("page_type"), p.get("legibility"),
         p.get("transcript"), "\n".join(p.get("unreadable_regions") or []), p.get("has_stamp_or_handwriting"),
         p.get("has_table")] for p in pages])

    issues = proposals.get("review_issues", [])
    _append_sheet(wb, "Extraction_Notes", [
        "Severity", "Code", "Message", "Field", "Page", "Related item",
    ], [[r.get("severity"), r.get("code"), r.get("message"), r.get("field"), r.get("page"), r.get("related_id")]
        for r in issues])

    changes = sorted((c for c in list_records("agreement_changes")
                      if c.get("addendum_document_id") == document["id"]), key=lambda record: record["id"])
    _append_sheet(wb, "Addendum_Changes", [
        "Change type", "Target", "Field", "Prior value", "Extracted value", "Amendment date",
        "Effective date", "Status", "Note", "Source page", "Source quote", "Evidence found",
    ], [[r.get("kind"), r.get("target"), r.get("field"), r.get("prior_value"), r.get("proposed_value"),
         r.get("amendment_date"), r.get("effective_date"), r.get("status"), r.get("note"),
         (r.get("evidence") or {}).get("page"), (r.get("evidence") or {}).get("quote"),
         (r.get("evidence") or {}).get("verified")] for r in changes])

    wb.active = 0
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
