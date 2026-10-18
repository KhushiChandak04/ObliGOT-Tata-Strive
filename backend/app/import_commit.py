"""Import a messy CSV/XLSX sheet: clean it deterministically, then save the valid rows as operational records.

No AI is involved. Headers are matched to ObliGOT fields by known synonyms; Indian amounts (₹, lakh, crore) and
common date formats (day first) are normalised; every row goes through the same Pydantic models as the forms.
Nothing is guessed: a value that cannot be read is reported, not invented. Imported records feed alerts like any
other record, because alerts are always computed from current records.
"""
import csv
import io
import re
from datetime import date, datetime, timedelta
from typing import Any

from fastapi import APIRouter, HTTPException, Request
from pydantic import ValidationError

from app.config import settings
from app.models import AgreementCreate, FinanceCreate, OpportunityCreate, PartnerCreate, ProjectCreate, ReportCreate
from app.repository import list_records, patch_record, put_record
from app.reviews import actor, log_activity
from app.rules import business_today

router = APIRouter(tags=["import"])

PARTNER_NAME = ["partner", "partner name", "organisation", "organization", "org", "organisation name", "organization name",
                "donor", "funder", "company", "company name", "csr partner", "client", "corporate", "funding partner"]
OWNER = ["owner", "relationship manager", "rm", "poc", "spoc", "manager", "assigned to", "responsible", "point of contact"]
STATUS = ["status", "current status", "stage", "state"]
START = ["start", "start date", "effective date", "effective from", "from", "commencement", "commencement date", "signed on"]
END = ["end", "end date", "expiry", "expiry date", "expires", "valid till", "valid until", "to", "termination date", "completion date"]
DUE = ["due", "due date", "deadline", "due on", "submission deadline", "due by"]

FIELDS: dict[str, dict[str, list[str]]] = {
    "partners": {
        "name": PARTNER_NAME + ["name"], "category": ["category", "type", "sector", "partner type", "segment"],
        "status": STATUS, "owner": OWNER, "contact_name": ["contact", "contact person", "contact name"],
        "contact_email": ["email", "e mail", "contact email", "mail", "email id"],
        "summary": ["notes", "remarks", "summary", "comments", "description"],
        "tier": ["tier", "customer tier", "partner tier", "category tier", "class", "membership", "level"],
        "group": ["group", "tata group", "tata non tata", "tata or non tata", "group company", "tata"],
        "location": ["location", "city", "state", "region", "hq", "headquarters"],
    },
    "agreements": {
        "partner": PARTNER_NAME, "title": ["title", "agreement", "mou", "mou title", "agreement name", "agreement title",
                                           "programme", "program", "project", "name"],
        "reference_number": ["ref", "reference", "reference number", "ref no", "mou no", "mou number", "agreement no"],
        "start_date": START, "expiry_date": END, "status": STATUS, "owner": OWNER,
        "signed_copy": ["signed copy", "signed", "signed copy received"],
    },
    "projects": {
        "partner": PARTNER_NAME, "name": ["project", "project name", "programme", "program", "name", "title"],
        "status": STATUS, "owner": OWNER, "start_date": START, "end_date": END,
        "milestone_summary": ["milestones", "milestone", "notes", "remarks"],
    },
    "finance": {
        "partner": PARTNER_NAME, "title": ["title", "description", "particulars", "item", "name", "details", "narration"],
        "record_type": ["record type", "type", "category", "finance type", "document type"],
        "amount": ["amount", "value", "amount inr", "amount rs", "amount in inr", "amount in rs", "sum", "total", "grant",
                   "grant amount", "inr", "rs"],
        "currency": ["currency", "ccy"], "issue_date": ["issue date", "invoice date", "date", "dated"],
        "due_date": DUE, "received_date": ["received", "received date", "received on", "payment date", "date received"],
        "status": STATUS, "reference_number": ["invoice no", "invoice number", "ref", "reference", "uc no", "reference number"],
        "owner": OWNER,
    },
    "opportunities": {
        "partner": PARTNER_NAME, "title": ["deal", "deal name", "opportunity", "opportunity name", "proposal", "title", "name",
                                           "programme", "program", "project"],
        "stage": STATUS + ["pipeline stage", "deal stage"], "owner": OWNER,
        "location": ["location", "city", "state", "geography", "region", "district"],
        "amount": ["amount", "deal size", "deal amount", "ask", "proposed amount", "value", "budget", "amount inr",
                   "amount in inr", "expected amount", "grant amount"],
        "sanction_probability": ["probability", "sanction probability", "likelihood", "chance", "probability of sanction",
                                 "win probability", "conversion probability"],
        "started_on": ["started", "started on", "start date", "since", "deal started", "initiated on", "first contact",
                       "date initiated", "start"],
        "last_followup_date": ["last follow up", "last followup", "follow up", "followup", "last contact",
                               "last follow up date", "follow up date", "last connected"],
        "next_step": ["next step", "next action", "action"], "next_step_date": ["next step date", "next action date"],
    },
    "reports": {
        "partner": PARTNER_NAME, "title": ["report", "title", "name", "deliverable", "report name", "report title"],
        "report_type": ["type", "report type", "frequency"], "due_date": DUE,
        "submitted_date": ["submitted", "submitted on", "submission date", "date submitted"], "status": STATUS,
        "owner": OWNER,
    },
}
MODELS = {"partners": PartnerCreate, "agreements": AgreementCreate, "projects": ProjectCreate,
          "finance": FinanceCreate, "reports": ReportCreate, "opportunities": OpportunityCreate}
# Partner fields a re-imported customer sheet may fill in on existing partners.
PARTNER_SHEET_FIELDS = ("tier", "group", "location")
DATE_FIELDS = {"start_date", "expiry_date", "end_date", "issue_date", "due_date", "received_date", "submitted_date",
               "started_on", "last_followup_date", "next_step_date"}

# Common spreadsheet words → the model's values. Unknown values are reported and the model default is kept.
ENUMS: dict[tuple[str, str], dict[str, str]] = {
    ("partners", "status"): {"active": "active", "live": "active", "ongoing": "active", "prospect": "prospect",
                             "lead": "prospect", "new": "prospect", "qualification": "qualification",
                             "in discussion": "qualification", "on hold": "on_hold", "hold": "on_hold", "paused": "on_hold",
                             "closed": "closed", "inactive": "closed", "dropped": "closed"},
    ("partners", "tier"): {"platinum": "platinum", "plat": "platinum", "gold": "gold", "silver": "silver",
                           "standard": "standard", "regular": "standard", "normal": "standard", "bronze": "standard"},
    ("partners", "group"): {"tata": "tata_group", "tata group": "tata_group", "yes": "tata_group", "y": "tata_group",
                            "group company": "tata_group", "non tata": "non_tata", "nontata": "non_tata", "no": "non_tata",
                            "n": "non_tata", "external": "non_tata", "others": "non_tata", "other": "non_tata"},
    ("agreements", "status"): {"active": "active", "signed": "active", "live": "active", "ongoing": "active",
                               "in force": "active", "draft": "draft", "under review": "under_review",
                               "in review": "under_review", "pending": "under_review", "expired": "expired",
                               "lapsed": "expired", "closed": "closed", "terminated": "closed", "completed": "closed"},
    ("agreements", "signed_copy"): {"yes": "received", "y": "received", "received": "received", "signed": "received",
                                    "no": "pending", "n": "pending", "pending": "pending", "unknown": "unknown"},
    ("projects", "status"): {"planned": "planned", "not started": "planned", "upcoming": "planned", "active": "active",
                             "ongoing": "active", "in progress": "active", "live": "active", "on hold": "on_hold",
                             "paused": "on_hold", "completed": "completed", "done": "completed", "closed": "closed"},
    ("finance", "record_type"): {"budget": "budget", "allocation": "budget", "sanction": "budget",
                                 "fund received": "fund_received", "funds received": "fund_received",
                                 "receipt": "fund_received", "payment received": "fund_received", "received": "fund_received",
                                 "disbursement": "fund_received", "invoice": "invoice", "inv": "invoice", "bill": "invoice",
                                 "uc": "utilization_certificate", "utilization certificate": "utilization_certificate",
                                 "utilisation certificate": "utilization_certificate"},
    ("opportunities", "stage"): {"identified": "identified", "lead": "identified", "new": "identified", "prospect": "identified",
                                 "qualification": "qualification", "qualified": "qualification", "in discussion": "qualification",
                                 "proposal": "proposal", "proposal sent": "proposal", "proposal submitted": "proposal",
                                 "approval": "approval", "approved": "approval", "sanctioned": "approval",
                                 "agreement": "agreement", "mou": "agreement", "mou signed": "agreement", "signed": "agreement",
                                 "active": "active", "ongoing": "active", "renewal": "renewal", "closed": "closed",
                                 "lost": "closed", "dropped": "closed", "won": "active"},
    ("reports", "report_type"): {"monthly": "monthly", "month": "monthly", "quarterly": "quarterly", "quarter": "quarterly",
                                 "q": "quarterly", "impact": "impact", "impact report": "impact"},
    ("reports", "status"): {"planned": "planned", "not started": "planned", "pending": "planned", "upcoming": "planned",
                            "due": "due", "submitted": "submitted", "sent": "submitted", "done": "submitted",
                            "completed": "submitted", "overdue": "overdue", "delayed": "overdue", "late": "overdue",
                            "accepted": "accepted", "approved": "accepted"},
}


def _norm(text: Any) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(text or "").lower())).strip()


def map_headers(headers: list[str], entity_type: str) -> dict[int, str]:
    """Column index → field. Exact synonym first, then the longest synonym the header starts with; each field once."""
    fields, mapping, used = FIELDS[entity_type], {}, set()
    normalised = [_norm(h) for h in headers]
    for exact in (True, False):
        for index, header in enumerate(normalised):
            if index in mapping or not header:
                continue
            best = None
            for field, synonyms in fields.items():
                if field in used:
                    continue
                for synonym in synonyms:
                    hit = header == synonym if exact else (header.startswith(synonym + " ") or header.endswith(" " + synonym))
                    if hit and (best is None or len(synonym) > best[1]):
                        best = (field, len(synonym))
            if best:
                mapping[index] = best[0]
                used.add(best[0])
    return mapping


_MULTIPLIERS = [(r"(crores?|cr)\b", 10_000_000), (r"(lakhs?|lacs?|lac|l)\b", 100_000), (r"(millions?|mn|m)\b", 1_000_000),
                (r"(thousand|k)\b", 1_000)]


def clean_amount(value: Any) -> tuple[float | None, str | None, str | None]:
    """(amount, currency if stated, error). '₹12,50,000', '12.5 lakh', '1.2 Cr', 'Rs. 5000/-' → numbers."""
    if value is None or str(value).strip() in ("", "-", "na", "n/a", "nil"):
        return None, None, None
    if isinstance(value, (int, float)):
        return (float(value), None, None) if value >= 0 else (None, None, "amount cannot be negative")
    text = str(value).strip().lower()
    currency = "INR" if re.search(r"₹|\brs\b|\binr\b|rs\.", text) else "USD" if re.search(r"\$|\busd\b", text) else \
        "EUR" if re.search(r"€|\beur\b", text) else None
    negative = text.startswith("-") or (text.startswith("(") and text.endswith(")"))
    text = re.sub(r"₹|\$|€|\brs\b\.?|\binr\b|\busd\b|\beur\b|/-|\(|\)|,|\s", " ", text).strip()
    multiplier = 1
    for pattern, factor in _MULTIPLIERS:
        if re.search(r"\d\s*" + pattern, text):
            multiplier, text = factor, re.sub(pattern, "", text)
            break
    text = text.replace(" ", "")
    try:
        amount = float(text) * multiplier
    except ValueError:
        return None, currency, f"amount “{value}” is not a number"
    if negative:
        return None, currency, "amount cannot be negative"
    return round(amount, 2), currency, None


_DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y", "%d.%m.%y", "%d %b %Y",
                 "%d %B %Y", "%d-%b-%Y", "%d-%B-%Y", "%d-%b-%y", "%d %b %y", "%b %d %Y", "%B %d %Y", "%Y/%m/%d",
                 "%b %Y", "%B %Y"]


def clean_date(value: Any) -> tuple[date | None, str | None]:
    """Day-first (Indian) dates in common formats, Excel serial numbers and real date cells."""
    if value is None or str(value).strip() in ("", "-", "na", "n/a", "tbd", "tba"):
        return None, None
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # Excel serial date
        return date(1899, 12, 30) + timedelta(days=int(value)), None
    text = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", str(value).strip(), flags=re.I).replace(",", " ")
    text = re.sub(r"\s+", " ", text.split("T")[0]).strip()
    text = re.sub(r"^(\d{4}-\d{2}-\d{2}) \d{1,2}:\d{2}(:\d{2})?$", r"\1", text)
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date(), None
        except ValueError:
            continue
    return None, f"date “{value}” could not be read"


def _text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def read_sheet(filename: str, raw: bytes) -> list[list[Any]]:
    """All rows of the file (first sheet for XLSX), cell values kept as typed by Excel."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        try:
            return [row for row in csv.reader(io.StringIO(raw.decode("utf-8-sig")))]
        except UnicodeDecodeError:
            return [row for row in csv.reader(io.StringIO(raw.decode("cp1252", errors="replace")))]
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            return [list(row) for row in workbook.active.iter_rows(values_only=True)]
        finally:
            workbook.close()
    raise HTTPException(status_code=415, detail="Upload a .csv or .xlsx file")


def find_header(rows: list[list[Any]], entity_type: str) -> tuple[int, dict[int, str]]:
    """Messy sheets often have a title or blank rows first: use the row (of the first 15) that maps the most fields."""
    best = (0, {})
    for index, row in enumerate(rows[:15]):
        mapping = map_headers([_text(c) for c in row], entity_type)
        if len(mapping) > len(best[1]):
            best = (index, mapping)
    return best


def clean_rows(entity_type: str, rows: list[list[Any]]) -> dict:
    header_index, mapping = find_header(rows, entity_type)
    headers = [_text(c) for c in rows[header_index]] if rows else []
    cleaned, rejected = [], []
    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(_text(c) for c in row):
            continue
        data: dict[str, Any] = {}
        errors, warnings = [], []
        stated_currency = None
        for index, field in mapping.items():
            value = row[index] if index < len(row) else None
            if field == "amount":
                amount, stated_currency, error = clean_amount(value)
                if error:
                    errors.append(error)
                data["amount"] = amount
            elif field == "sanction_probability":
                text = _text(value).replace("%", "").strip()
                if text:
                    try:
                        number = float(text)
                        number = number * 100 if 0 < number < 1 else number  # 0.6 written as a fraction
                        if 0 <= number <= 100:
                            data[field] = round(number)
                        else:
                            errors.append(f"probability “{value}” is not between 0 and 100")
                    except ValueError:
                        warnings.append(f"probability “{value}” is not a number; left blank")
            elif field in DATE_FIELDS:
                parsed, error = clean_date(value)
                if error:
                    errors.append(error)
                data[field] = parsed.isoformat() if parsed else None
            else:
                text = _text(value)
                if (entity_type, field) in ENUMS and text:
                    known = ENUMS[(entity_type, field)].get(_norm(text))
                    if known:
                        data[field] = known
                    else:
                        warnings.append(f"{field.replace('_', ' ')} “{text}” not recognised; default used")
                elif field == "currency":
                    word = _norm(text)
                    data[field] = "INR" if word in {"rs", "rupees", "rupee", "inr"} or "₹" in text else                         "USD" if word in {"usd", "dollar", "dollars"} or "$" in text else text.upper()[:3] or None
                else:
                    data[field] = text or None
        if stated_currency and not data.get("currency"):
            data["currency"] = stated_currency
        data = {k: v for k, v in data.items() if v not in (None, "")}
        item = {"row": offset, "data": data, "warnings": warnings, "errors": errors}
        (rejected if errors else cleaned).append(item)
    return {"header_row": header_index + 1, "headers": headers,
            "mapping": {headers[i] or f"Column {i + 1}": f for i, f in sorted(mapping.items())},
            "unmapped": [h for i, h in enumerate(headers) if h and i not in mapping], "cleaned": cleaned, "rejected": rejected}


def _key(entity_type: str, d: dict) -> tuple:
    if entity_type == "partners":
        return (_norm(d.get("name")),)
    if entity_type == "projects":
        return (d.get("partner_id"), _norm(d.get("name")))
    if entity_type == "finance":
        return (d.get("partner_id"), d.get("record_type"), _norm(d.get("title")), d.get("amount"), d.get("due_date"))
    if entity_type == "reports":
        return (d.get("partner_id"), _norm(d.get("title")), d.get("due_date"))
    if entity_type == "opportunities":
        return (d.get("partner_id"), _norm(d.get("title")))
    return (d.get("partner_id"), _norm(d.get("title")))


@router.post("/import/commit")
async def import_commit(request: Request) -> dict:
    """Clean an uploaded sheet and, unless dry_run, save the valid rows. Duplicates of existing records are skipped."""
    if "multipart/form-data" not in request.headers.get("content-type", "").lower():
        raise HTTPException(status_code=415, detail="Upload a CSV/XLSX file as multipart/form-data")
    form = await request.form()
    entity_type = str(form.get("entity_type") or "projects")
    entity_type = "finance" if entity_type == "finance-records" else entity_type
    if entity_type not in MODELS:
        raise HTTPException(status_code=422, detail="Unsupported entity_type")
    dry_run = str(form.get("dry_run") or "").lower() == "true"
    upload = form.get("file")
    if upload is None or not hasattr(upload, "read"):
        raise HTTPException(status_code=422, detail="A file field is required")
    raw = await upload.read()
    if len(raw) > 5_000_000:
        raise HTTPException(status_code=413, detail="File is larger than the 5 MB import limit")
    try:
        rows = read_sheet(getattr(upload, "filename", "") or "", raw)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Could not read this file; check its format") from exc
    if len(rows) > 2001:
        raise HTTPException(status_code=413, detail="Import at most 2,000 rows at a time")
    result = clean_rows(entity_type, rows)
    if entity_type != "partners" and "partner" not in result["mapping"].values():
        raise HTTPException(status_code=422, detail="No partner column found; add a column such as “Partner” or “Organisation”")

    partners = {_norm(p["name"]): p for p in list_records("partners")}
    existing = {_key(entity_type, r) for r in list_records(entity_type)}
    new_partners: dict[str, dict] = {}
    to_save, skipped, to_update = [], [], []
    model = MODELS[entity_type]
    for item in result["cleaned"]:
        data = dict(item["data"])
        if entity_type == "partners":
            data.setdefault("category", "Not specified")
            data.setdefault("owner", "Unassigned")
            if "status" not in data:
                item["warnings"].append("status not in the file; saved as prospect")
        else:
            name = data.pop("partner", None)
            if not name:
                item["errors"].append("partner name is missing")
                continue
            partner = partners.get(_norm(name)) or new_partners.get(_norm(name))
            if partner is None:
                partner = {"id": f"partner-new-{len(new_partners) + 1}", "name": name}
                new_partners[_norm(name)] = partner
                item["warnings"].append(f"new partner “{name}” will be created")
            data["partner_id"] = partner["id"]
            if entity_type == "agreements" and "status" not in data:
                item["warnings"].append("status not in the file; saved as draft — set it to active to get expiry alerts")
            if entity_type == "finance" and "record_type" not in data:
                item["errors"].append("finance type is missing (budget, fund received, invoice or UC)")
                continue
        try:
            validated = model(**data).model_dump(mode="json")
        except ValidationError as exc:
            item["errors"].extend(f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors())
            continue
        key = _key(entity_type, validated)
        if key in existing:
            # A customer sheet often lists partners already in ObliGOT: fill in tier, group and location from it.
            current = partners.get(key[0]) if entity_type == "partners" else None
            changes = {k: item["data"][k] for k in PARTNER_SHEET_FIELDS if current and item["data"].get(k) not in (None, "")
                       and item["data"][k] != current.get(k)}
            if changes:
                to_update.append((current, changes))
            skipped.append({"row": item["row"], "reason": "already exists; tier/group/location updated" if changes else "already exists"})
            continue
        existing.add(key)
        to_save.append((item, validated))
    result["rejected"] += [i for i in result["cleaned"] if i["errors"]]
    result["rejected"].sort(key=lambda i: i["row"])
    skipped_rows = {s["row"] for s in skipped}
    result["cleaned"] = [i for i in result["cleaned"] if not i["errors"] and i["row"] not in skipped_rows]

    created = 0
    if not dry_run:
        today = business_today().isoformat()
        for current, changes in to_update:
            patch_record("partners", current["id"], {**changes, "updated_at": today})
        for placeholder in new_partners.values():
            saved = put_record("partners", {**PartnerCreate(name=placeholder["name"], category="Not specified", status="active",
                                                            owner="Unassigned", summary="Created by spreadsheet import.")
                                            .model_dump(mode="json"), "updated_at": today, "source_label": "user"})
            placeholder["saved_id"] = saved["id"]
        for item, validated in to_save:
            pid = validated.get("partner_id")
            for placeholder in new_partners.values():
                if pid == placeholder["id"]:
                    validated["partner_id"] = placeholder["saved_id"]
            extra = {"updated_at": today} if entity_type in ("partners", "opportunities") else {}
            put_record(entity_type, {**validated, **extra, "source_label": "user", "imported": True})
            created += 1
        filename = getattr(upload, "filename", "") or "sheet"
        log_activity(actor_name=actor(request), event="records_imported",
                     summary=f"Imported {created} {entity_type} from “{filename}” ({len(result['rejected'])} rejected, "
                             f"{len(skipped)} duplicates skipped, {len(new_partners)} new partners)",
                     entity_type="import", entity_id=f"import-{entity_type}", partner_id=None,
                     details={"created": created, "rejected": len(result["rejected"]), "skipped": len(skipped)})
    return {**result, "entity_type": entity_type, "dry_run": dry_run, "committed": not dry_run,
            "valid_count": len(to_save), "created": created, "skipped": skipped, "updated": len(to_update),
            "new_partners": [p["name"] for p in new_partners.values()], "default_currency": settings.default_currency}
